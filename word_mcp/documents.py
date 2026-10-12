from __future__ import annotations

import base64
import hashlib
import io
import json
import os
import re
import tempfile
import uuid
import zipfile
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph
from PIL import Image, ImageOps

from .config import Settings
from .schemas import (
    AddParagraph, AddPicture, AddTable, EditPlan, FormatParagraph,
    ReplacePicture, ReplaceText,
)


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


class DocumentStore:
    """File-based DOCX operations. Edits are atomic and guarded by a revision hash."""

    def __init__(self, settings: Settings):
        self.settings = settings
        self.root = settings.root.resolve()
        self.root.mkdir(parents=True, exist_ok=True)
        self.internal = self.root / ".word_mcp"
        if self.internal.is_symlink() or not self.internal.resolve().is_relative_to(self.root):
            raise ValueError("Internal storage must be inside the document directory")
        for subdir in ("backups", "proposals", "locks"):
            directory = self.internal / subdir
            if directory.is_symlink():
                raise ValueError("Internal storage cannot contain symlinks")
            directory.mkdir(parents=True, exist_ok=True)

    def path(self, name: str, suffix: str | None = None) -> Path:
        if not name or "\x00" in name or "\\" in name:
            raise ValueError("Use a relative path with forward slashes")
        relative = Path(name)
        if relative.is_absolute() or ":" in name or ".." in relative.parts:
            raise ValueError("Only relative paths inside WORD_MCP_ROOT are allowed")
        if any(part.startswith(".") for part in relative.parts):
            raise ValueError("Hidden/internal paths are not accessible")
        target = (self.root / relative).resolve()
        if not target.is_relative_to(self.root):
            raise ValueError("Path escapes the document directory")
        if suffix and target.suffix.lower() != suffix:
            raise ValueError(f"Only {suffix} files are supported")
        return target

    def validate_docx(self, data: bytes) -> None:
        if len(data) > self.settings.max_file_bytes:
            raise ValueError("Document exceeds the 30 MiB limit")
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                infos = archive.infolist()
                if len(infos) > 10000 or sum(x.file_size for x in infos) > self.settings.max_unpacked_bytes:
                    raise ValueError("DOCX archive is too large when unpacked")
                if "word/document.xml" not in archive.namelist():
                    raise ValueError("This is not a DOCX document")
                if any(x.flag_bits & 1 for x in infos):
                    raise ValueError("Encrypted documents are unsupported")
                if any("vbaproject" in x.filename.lower() for x in infos):
                    raise ValueError("Macro-enabled documents are unsupported")
            Document(io.BytesIO(data))
        except (zipfile.BadZipFile, KeyError) as exc:
            raise ValueError("Invalid DOCX archive") from exc

    def load(self, name: str):
        path = self.path(name, ".docx")
        if path.stat().st_size > self.settings.max_file_bytes:
            raise ValueError("Document exceeds the 30 MiB limit")
        data = path.read_bytes()
        self.validate_docx(data)
        return path, data, Document(io.BytesIO(data))

    @contextmanager
    def lock(self, path: Path):
        lockfile = self.internal / "locks" / (sha256(str(path).encode()) + ".lock")
        try:
            descriptor = os.open(lockfile, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
        except FileExistsError as exc:
            raise ValueError("Another operation is editing this document. Retry after it finishes.") from exc
        try:
            os.write(descriptor, str(os.getpid()).encode())
            os.close(descriptor)
            yield
        finally:
            lockfile.unlink(missing_ok=True)

    @staticmethod
    def atomic_write(path: Path, data: bytes):
        path.parent.mkdir(parents=True, exist_ok=True)
        descriptor, temporary = tempfile.mkstemp(prefix=".word-write-", dir=path.parent)
        try:
            with os.fdopen(descriptor, "wb") as output:
                output.write(data)
                output.flush()
                os.fsync(output.fileno())
            os.replace(temporary, path)
        finally:
            Path(temporary).unlink(missing_ok=True)

    def list_documents(self) -> list[dict]:
        results = []
        for path in sorted(self.root.rglob("*")):
            if path.suffix.lower() != ".docx" or not path.is_file():
                continue
            name = path.relative_to(self.root).as_posix()
            try:
                self.path(name, ".docx")
            except ValueError:
                continue
            results.append({"path": name, "bytes": path.stat().st_size})
            if len(results) >= 500:
                break
        return results

    def import_document(self, name: str, data: bytes) -> dict:
        path = self.path(name, ".docx")
        self.validate_docx(data)
        with self.lock(path):
            if path.exists():
                raise ValueError("A document with this name already exists")
            self.atomic_write(path, data)
        return {"path": name, "revision": sha256(data)}

    def create(self, name: str, title: str = "") -> dict:
        document = Document()
        if title:
            document.add_paragraph(title, "Title")
        else:
            document.add_paragraph()
        output = io.BytesIO()
        document.save(output)
        return self.import_document(name, output.getvalue())

    @staticmethod
    def paragraphs(document) -> dict[str, Paragraph]:
        # Includes table cells and nested tables, then headers and footers.
        result = {f"p{i}": Paragraph(element, document)
                  for i, element in enumerate(document.element.body.iter(qn("w:p")))}
        for part in sorted(document.part.package.parts, key=lambda x: str(x.partname)):
            match = re.fullmatch(r"/word/(header|footer)(\d+)\.xml", str(part.partname))
            if match and hasattr(part, "element"):
                prefix = match.group(1) + match.group(2)
                for i, element in enumerate(part.element.iter(qn("w:p"))):
                    result[f"{prefix}:p{i}"] = Paragraph(element, document)
        return result

    @staticmethod
    def image_parts(document) -> dict:
        return {str(part.partname).lstrip("/"): part
                for part in document.part.package.parts
                if str(part.content_type).startswith("image/") and str(part.partname).startswith("/word/")}

    def snapshot(self, name: str) -> dict:
        _, data, document = self.load(name)
        paragraphs = [
            {"id": key, "text": paragraph.text, "style": paragraph.style.name,
             "has_complex_content": self.is_complex(paragraph)}
            for key, paragraph in self.paragraphs(document).items()
        ]
        images = []
        for key, part in self.image_parts(document).items():
            info = {"id": key, "content_type": part.content_type, "bytes": len(part.blob)}
            try:
                with Image.open(io.BytesIO(part.blob)) as image:
                    info.update(width=image.width, height=image.height)
            except (OSError, ValueError):
                info["preview_supported"] = False
            images.append(info)
        tables = []
        for i, table in enumerate(document.tables):
            tables.append({"id": f"t{i}", "rows": [[cell.text for cell in row.cells] for row in table.rows]})
        return {"path": name, "revision": sha256(data), "paragraphs": paragraphs,
                "tables": tables, "images": images,
                "notes": ["Paragraph IDs are valid only with this revision.",
                          "Text in fields, hyperlinks, text boxes and tracked changes is read-only.",
                          "This is a structure preview, not a page layout rendering."]}

    @staticmethod
    def is_complex(paragraph: Paragraph) -> bool:
        protected = {"hyperlink", "fldSimple", "fldChar", "instrText", "ins", "del", "sdt", "txbxContent", "oMath"}
        return any(element.tag.rsplit("}", 1)[-1] in protected for element in paragraph._p.iter())

    @staticmethod
    def replace_text(paragraph: Paragraph, old: str, new: str, occurrence: int):
        if DocumentStore.is_complex(paragraph):
            raise ValueError("This paragraph contains fields, hyperlinks or tracked changes; editing it is unsupported")
        runs = paragraph.runs
        combined = "".join(run.text for run in runs)
        matches = list(re.finditer(re.escape(old), combined))
        if occurrence >= len(matches):
            raise ValueError(f"Expected text {old!r} was not found at occurrence {occurrence}")
        begin, end = matches[occurrence].span()
        cursor = 0
        spans = []
        for run in runs:
            start, stop = cursor, cursor + len(run.text)
            if start < end and stop > begin:
                # Assigning run.text would remove drawings/breaks in that run.
                if any(child.tag not in (qn("w:rPr"), qn("w:t")) for child in run._r):
                    raise ValueError("The text overlaps a run containing a drawing, field or break")
                spans.append((run, start, stop, run.text))
            cursor = stop
        for i, (run, start, stop, text) in enumerate(spans):
            before = text[:max(0, begin - start)]
            after = text[max(0, end - start):] if end < stop else ""
            run.text = before + (new if i == 0 else "") + after

    @staticmethod
    def image_bytes(data: bytes, edge: int | None = None) -> bytes:
        if len(data) > 20 * 1024 * 1024:
            raise ValueError("Image exceeds 20 MiB")
        try:
            with Image.open(io.BytesIO(data)) as source:
                if source.width * source.height > 25_000_000:
                    raise ValueError("Image exceeds 25 megapixels")
                image = ImageOps.exif_transpose(source).convert("RGB")
                if edge:
                    image.thumbnail((edge, edge))
                output = io.BytesIO()
                image.save(output, format="PNG")
                return output.getvalue()
        except OSError as exc:
            raise ValueError("Unsupported raster image. Use PNG, JPEG or WebP") from exc

    def save_asset(self, filename: str, data: bytes) -> dict:
        normalized = self.image_bytes(data)
        name = f"assets/{uuid.uuid4().hex}.png"
        self.atomic_write(self.path(name), normalized)
        return {"path": name, "original_name": Path(filename).name}

    def image(self, name: str, image_id: str, thumbnail: bool = False) -> bytes:
        _, _, document = self.load(name)
        parts = self.image_parts(document)
        if image_id not in parts:
            raise ValueError("Image ID not found. Read the document first")
        return self.image_bytes(parts[image_id].blob, 768 if thumbnail else None)

    def apply_operations(self, document, plan: EditPlan) -> list[dict]:
        paragraphs = self.paragraphs(document)  # Stable IDs, even when insertions occur.
        parts = self.image_parts(document)
        changes = []
        added_after = {}
        for operation in plan.operations:
            def get_paragraph(target):
                if target not in paragraphs:
                    raise ValueError(f"Unknown paragraph ID: {target}")
                return paragraphs[target]

            def append_paragraph(text, style, after):
                if after is None:
                    return document.add_paragraph(text, style)
                anchor = added_after.get(after)
                if anchor is None:
                    anchor = get_paragraph(after)
                if self.is_complex(anchor):
                    raise ValueError("Cannot insert into complex paragraph content")
                element = OxmlElement("w:p")
                anchor._p.addnext(element)
                paragraph = Paragraph(element, anchor._parent)
                paragraph.style = style
                paragraph.add_run(text)
                added_after[after] = paragraph
                return paragraph

            if isinstance(operation, ReplaceText):
                paragraph = get_paragraph(operation.target)
                before = paragraph.text
                self.replace_text(paragraph, operation.old, operation.new, operation.occurrence)
                changes.append({"op": operation.op, "target": operation.target, "before": before, "after": paragraph.text})
            elif isinstance(operation, AddParagraph):
                append_paragraph(operation.text, operation.style, operation.after)
                changes.append(operation.model_dump())
            elif isinstance(operation, FormatParagraph):
                paragraph = get_paragraph(operation.target)
                if self.is_complex(paragraph):
                    raise ValueError("Formatting complex paragraphs is unsupported")
                for run in paragraph.runs:
                    if operation.bold is not None:
                        run.bold = operation.bold
                    if operation.italic is not None:
                        run.italic = operation.italic
                    if operation.size_pt is not None:
                        run.font.size = Pt(operation.size_pt)
                changes.append(operation.model_dump())
            elif isinstance(operation, AddTable):
                columns = len(operation.rows[0])
                if not 1 <= columns <= 30 or any(len(row) != columns for row in operation.rows):
                    raise ValueError("A table must have 1–30 columns and equal row lengths")
                if sum(len(cell) for row in operation.rows for cell in row) > 100000:
                    raise ValueError("Table content is too large")
                table = document.add_table(rows=len(operation.rows), cols=columns)
                table.style = "Table Grid"
                for row, values in zip(table.rows, operation.rows):
                    for cell, value in zip(row.cells, values):
                        cell.text = value
                changes.append(operation.model_dump())
            elif isinstance(operation, AddPicture):
                image_path = self.path(operation.image_path)
                blob = self.image_bytes(image_path.read_bytes())
                paragraph = append_paragraph("", "Normal", operation.after)
                shape = paragraph.add_run().add_picture(io.BytesIO(blob), width=Cm(operation.width_cm))
                shape._inline.docPr.set("descr", operation.alt_text)
                changes.append(operation.model_dump())
            elif isinstance(operation, ReplacePicture):
                if operation.image_id not in parts:
                    raise ValueError("Unknown image ID")
                part = parts[operation.image_id]
                formats = {"image/png": "PNG", "image/jpeg": "JPEG", "image/gif": "GIF",
                           "image/bmp": "BMP", "image/tiff": "TIFF"}
                if part.content_type not in formats:
                    raise ValueError("Replacing this image format is unsupported. Add a PNG/JPEG picture instead")
                normalized = self.image_bytes(self.path(operation.image_path).read_bytes())
                with Image.open(io.BytesIO(normalized)) as replacement:
                    output = io.BytesIO()
                    replacement.save(output, format=formats[part.content_type])
                    part._blob = output.getvalue()
                changes.append({**operation.model_dump(), "note": "All uses of this shared image part are replaced; existing display size is preserved"})
        return changes

    def preview(self, name: str, revision: str, plan: EditPlan) -> dict:
        _, data, document = self.load(name)
        if sha256(data) != revision:
            raise ValueError("Document changed. Read it again and generate a new edit plan")
        changes = self.apply_operations(document, plan)
        return {"path": name, "revision": revision, "summary": plan.summary, "changes": changes}

    def edit(self, name: str, revision: str, plan: EditPlan) -> dict:
        path = self.path(name, ".docx")
        with self.lock(path):
            _, original, document = self.load(name)
            if sha256(original) != revision:
                raise ValueError("Document changed. Read it again and generate a new edit plan")
            changes = self.apply_operations(document, plan)
            output = io.BytesIO()
            document.save(output)
            edited = output.getvalue()
            self.validate_docx(edited)
            # Recheck before commit to catch edits made outside this application.
            if sha256(path.read_bytes()) != revision:
                raise ValueError("Document changed while the edit was being prepared")
            backup_id = uuid.uuid4().hex
            self.atomic_write(self.internal / "backups" / f"{backup_id}.docx", original)
            self.atomic_write(self.internal / "backups" / f"{backup_id}.json", json.dumps(
                {"path": name, "revision": revision, "time": datetime.now(timezone.utc).isoformat()},
                ensure_ascii=False).encode())
            self.atomic_write(path, edited)
        return {"path": name, "revision": sha256(edited), "backup_id": backup_id,
                "summary": plan.summary, "changes": changes}

    @staticmethod
    def checked_id(value: str) -> str:
        if not re.fullmatch(r"[0-9a-f]{32}", value):
            raise ValueError("Invalid ID")
        return value

    def backups(self, name: str) -> list[dict]:
        self.path(name, ".docx")
        result = []
        for path in (self.internal / "backups").glob("*.json"):
            record = json.loads(path.read_text(encoding="utf-8"))
            if record["path"] == name:
                result.append({"id": path.stem, **record})
        return sorted(result, key=lambda x: x["time"], reverse=True)

    def restore(self, name: str, backup_id: str, revision: str) -> dict:
        backup_id = self.checked_id(backup_id)
        record = json.loads((self.internal / "backups" / f"{backup_id}.json").read_text(encoding="utf-8"))
        if record["path"] != name:
            raise ValueError("Backup belongs to a different document")
        data = (self.internal / "backups" / f"{backup_id}.docx").read_bytes()
        self.validate_docx(data)
        path = self.path(name, ".docx")
        with self.lock(path):
            current = path.read_bytes()
            if sha256(current) != revision:
                raise ValueError("Document changed; reload it before restoring")
            undo_id = uuid.uuid4().hex
            self.atomic_write(self.internal / "backups" / f"{undo_id}.docx", current)
            self.atomic_write(self.internal / "backups" / f"{undo_id}.json", json.dumps(
                {"path": name, "revision": revision, "time": datetime.now(timezone.utc).isoformat()},
                ensure_ascii=False).encode())
            self.atomic_write(path, data)
        return {"path": name, "revision": sha256(data), "backup_id": undo_id}

    def save_proposal(self, name: str, revision: str, plan: EditPlan, model: str) -> dict:
        preview = self.preview(name, revision, plan)
        proposal_id = uuid.uuid4().hex
        # Freeze referenced assets so the approved preview cannot change later.
        for operation in plan.operations:
            if isinstance(operation, (AddPicture, ReplacePicture)):
                asset = self.save_asset("proposal.png", self.path(operation.image_path).read_bytes())
                operation.image_path = asset["path"]
        record = {"path": name, "revision": revision, "model": model, "plan": plan.model_dump()}
        self.atomic_write(self.internal / "proposals" / f"{proposal_id}.json", json.dumps(record, ensure_ascii=False).encode())
        return {"proposal_id": proposal_id, "model": model, **preview}

    def apply_proposal(self, proposal_id: str) -> dict:
        proposal_id = self.checked_id(proposal_id)
        path = self.internal / "proposals" / f"{proposal_id}.json"
        record = json.loads(path.read_text(encoding="utf-8"))
        result = self.edit(record["path"], record["revision"], EditPlan.model_validate(record["plan"]))
        path.unlink()
        return result
