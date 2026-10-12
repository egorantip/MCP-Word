"""Exercise the relocated portable folder with real local models and MCP."""
from __future__ import annotations

import argparse
import asyncio
from datetime import datetime, timezone
import io
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import time
import uuid

import httpx
from docx import Document
from docx.shared import Cm
from mcp import ClientSession, StdioServerParameters
from mcp.client.stdio import stdio_client
from mcp.client.streamable_http import streamablehttp_client
from PIL import Image, ImageDraw, ImageFont

from word_mcp.config import Settings
from word_mcp.documents import DocumentStore
from word_mcp.ollama_client import OllamaClient


ROOT = Path(__file__).resolve().parent.parent
REPORT = ROOT / "docs" / "verification.json"
FLAGS = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0


def wait_ready(url, process):
    deadline = time.monotonic() + 90
    with httpx.Client(timeout=2, trust_env=False) as client:
        while time.monotonic() < deadline:
            if process.poll() is not None:
                raise RuntimeError("Service exited before it was ready")
            try:
                if client.get(url).status_code == 200:
                    return
            except httpx.HTTPError:
                pass
            time.sleep(0.5)
    raise RuntimeError("Service startup timeout")


async def stdio_checks(env, name, models):
    params = StdioServerParameters(command=sys.executable, args=["-m", "word_mcp", "server"], env=env)
    outcomes = []
    async with stdio_client(params) as (read, write):
        async with ClientSession(read, write) as session:
            await session.initialize()
            tools = await session.list_tools()
            assert len(tools.tools) == 12

            async def call(tool, args):
                result = await session.call_tool(tool, args)
                if result.isError:
                    raise RuntimeError(result.content[0].text)
                return result.structuredContent or json.loads(result.content[0].text)

            for model in models:
                original = await call("read_document", {"path": name})
                print("Text test:", model, flush=True)
                began = time.monotonic()
                proposal = await call("propose_edit", {"path": name, "targets": ["p1"], "model": model,
                    "instruction": "В слове Превет опечатка. Замени только Превет на Привет. Больше ничего не меняй."})
                unchanged = await call("read_document", {"path": name})
                assert unchanged["revision"] == original["revision"]
                applied = await call("apply_proposal", {"proposal_id": proposal["proposal_id"]})
                result = await call("read_document", {"path": name})
                assert result["paragraphs"][1]["text"] == "Привет, мир."
                restored = await call("restore_backup", {"path": name, "backup_id": applied["backup_id"], "revision": result["revision"]})
                assert restored["revision"] == original["revision"]
                outcomes.append({"model": model, "text_edit": "passed", "preview": "passed", "restore": "passed",
                                 "seconds": round(time.monotonic() - began, 2), "proposal": proposal["changes"]})
    return outcomes


async def http_mcp_check(url, token):
    async with streamablehttp_client(url, headers={"Authorization": "Bearer " + token}) as (read, write, _):
        async with ClientSession(read, write) as session:
            await session.initialize()
            tools = await session.list_tools()
            assert len(tools.tools) == 12
            return "passed"


def child_checks():
    config = json.loads((ROOT / "portable.json").read_text(encoding="utf-8"))
    work = ROOT / "portable-qa"
    documents = work / uuid.uuid4().hex / "documents"
    work.mkdir(exist_ok=True)
    logs = ROOT / "logs"
    logs.mkdir(exist_ok=True)
    env = dict(os.environ)
    env.update({"WORD_MCP_ROOT": str(documents), "WORD_MCP_PROFILE": "light", "PYTHONUTF8": "1",
        "OLLAMA_HOST": "127.0.0.1:11635", "OLLAMA_MODELS": str(ROOT / "models"), "OLLAMA_NO_CLOUD": "1",
        "OLLAMA_MAX_LOADED_MODELS": "1", "OLLAMA_NUM_PARALLEL": "1", "OLLAMA_KEEP_ALIVE": "0",
        "CUDA_VISIBLE_DEVICES": "-1", "ROCR_VISIBLE_DEVICES": "-1", "GGML_VK_VISIBLE_DEVICES": "-1",
        "OLLAMA_VULKAN": "0", "USERPROFILE": str(ROOT / "runtime" / "user"),
        "TEMP": str(ROOT / "runtime" / "tmp"), "TMP": str(ROOT / "runtime" / "tmp")})
    Path(env["USERPROFILE"]).mkdir(parents=True, exist_ok=True)
    Path(env["TEMP"]).mkdir(parents=True, exist_ok=True)
    services = []
    files = []
    report = {"time": datetime.now(timezone.utc).isoformat(), "portable_root": str(ROOT), "python": sys.version,
              "internet_required_for_checks": False, "models": [], "status": "failed"}
    try:
        for executable, args, name in [
            (str(ROOT / "runtime" / "ollama" / "ollama.exe"), ["serve"], "verify-ollama"),
            (sys.executable, ["-m", "word_mcp", "web", "--port", "8876"], "verify-web")]:
            output = (logs / f"{name}.log").open("wb")
            files.append(output)
            services.append(subprocess.Popen([executable, *args], cwd=ROOT, env=env, creationflags=FLAGS,
                                            stdout=output, stderr=subprocess.STDOUT))
        wait_ready("http://127.0.0.1:11635/api/version", services[0])
        wait_ready("http://127.0.0.1:8876/health", services[1])
        settings = Settings(root=documents, ollama_url="http://127.0.0.1:11635", timeout=300)
        store = DocumentStore(settings)
        ollama = OllamaClient(settings)
        image = Image.new("RGB", (640, 320), "white")
        draw = ImageDraw.Draw(image)
        draw.rectangle((25, 25, 615, 170), fill="red")
        draw.text((50, 205), "MCP WORD 42", fill="black", font=ImageFont.load_default(size=50))
        picture = io.BytesIO()
        image.save(picture, format="PNG")
        doc = Document()
        doc.add_paragraph("Проверка MCP Word", "Title")
        paragraph = doc.add_paragraph()
        paragraph.add_run("Превет, ").bold = True
        paragraph.add_run("мир.").italic = True
        doc.add_picture(io.BytesIO(picture.getvalue()), width=Cm(10))
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Количество"
        table.cell(0, 1).text = "42"
        buffer = io.BytesIO()
        doc.save(buffer)
        name = "Живая проверка.docx"
        store.import_document(name, buffer.getvalue())
        text_models = [model for model in config["models"] if model.startswith("qwen3.5:")]
        report["models"] = asyncio.run(stdio_checks(env, name, text_models))
        snapshot = store.snapshot(name)
        image_id = snapshot["images"][0]["id"]
        for model in [settings.image_model]:
            outcome = {"model": model}
            report["models"].append(outcome)
            print("Vision test:", model, flush=True)
            start = time.monotonic()
            result = ollama.analyze(store, name, image_id,
                "Прочитай крупную надпись в нижней части картинки дословно. Укажи цвет прямоугольника. Предложи краткую подпись.", model)
            assert "42" in result["visible_text"], result
            assert result["description"].strip(), result
            outcome.update(vision="passed", vision_seconds=round(time.monotonic() - start, 2), image_analysis=result)
        token = (documents / ".word_mcp" / "http-token").read_text()
        headers = {"X-Word-Token": token}
        with httpx.Client(base_url="http://127.0.0.1:8876", headers=headers, trust_env=False) as client:
            assert client.get("/").status_code == 200
            read = client.get("/api/read", params={"path": name})
            assert read.status_code == 200
            downloaded = client.get("/download", params={"path": name})
            assert downloaded.content == buffer.getvalue()
            assert client.get("/image", params={"path": name, "id": image_id}).status_code == 200
            assert client.get("/health", headers={"Host": "evil.example"}).status_code == 403
        report["http_mcp"] = asyncio.run(http_mcp_check("http://127.0.0.1:8876/mcp/", token))
        assert not ollama.request("GET", "/api/ps").get("models"), "Model was not unloaded"
        assert Document(store.path(name)).paragraphs[1].runs[0].bold
        assert Document(store.path(name)).tables[0].cell(0, 1).text == "42"
        report.update(status="passed", http_ui="passed", download="passed", image_preview="passed",
                      model_unloading="passed", formatting="passed", actual_ollama_version=ollama.request("GET", "/api/version"))
    except Exception as exc:
        report["error"] = str(exc)
        raise
    finally:
        for service in reversed(services):
            service.terminate()
            try:
                service.wait(timeout=15)
            except subprocess.TimeoutExpired:
                service.kill()
        for file in files:
            file.close()
        REPORT.parent.mkdir(exist_ok=True)
        REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--child", action="store_true")
    parser.add_argument("--reuse-copy", action="store_true")
    args = parser.parse_args()
    if args.child:
        child_checks()
        return
    moved = ROOT / "portable-qa" / "Перенос" / "MCP Word"
    if moved.exists() and not args.reuse_copy:
        raise ValueError("The relocation test directory already exists; choose a fresh copy before repeating the test")
    print("Copying the prepared folder to a new Unicode path for relocation checks...", flush=True)
    moved.mkdir(parents=True, exist_ok=args.reuse_copy)
    for name in ["word_mcp", "scripts", "portable.json", "requirements.lock.txt", "Start.vbs", "Start.bat"]:
        source = ROOT / name
        if source.is_dir():
            shutil.copytree(source, moved / name, ignore=shutil.ignore_patterns("__pycache__"), dirs_exist_ok=args.reuse_copy)
        else:
            shutil.copy2(source, moved / name)
    for name in ([] if args.reuse_copy else ["python", "site-packages", "bootstrap", "ollama"]):
        print("Copying runtime/" + name, flush=True)
        shutil.copytree(ROOT / "runtime" / name, moved / "runtime" / name, ignore=shutil.ignore_patterns("__pycache__"))
    print("Copying all model caches...", flush=True)
    # Model blobs are immutable: when reusing a test copy, copy only new files.
    if args.reuse_copy:
        for source in (ROOT / "models").rglob("*"):
            if not source.is_file():
                continue
            destination = moved / "models" / source.relative_to(ROOT / "models")
            if not destination.exists() or destination.stat().st_size != source.stat().st_size:
                destination.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(source, destination)
    else:
        shutil.copytree(ROOT / "models", moved / "models")
    shutil.copy2(ROOT / "runtime" / "prepared.json", moved / "runtime" / "prepared.json")
    executable = moved / "runtime" / "python" / "python.exe"
    result = subprocess.run([str(executable), str(moved / "scripts" / "verify-live.py"), "--child"], cwd=moved,
                            creationflags=FLAGS)
    REPORT.parent.mkdir(exist_ok=True)
    shutil.copy2(moved / "docs" / "verification.json", REPORT)
    if result.returncode:
        raise SystemExit(result.returncode)
    print("All live checks passed. Report:", REPORT, flush=True)


if __name__ == "__main__":
    main()
