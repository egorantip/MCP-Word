import io
import zipfile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Cm
from PIL import Image

from word_mcp.config import Settings
from word_mcp.documents import DocumentStore, sha256
from word_mcp.schemas import EditPlan


@pytest.fixture
def store(tmp_path):
    return DocumentStore(Settings(root=tmp_path))


def plan(*operations):
    return EditPlan.model_validate({"summary": "Проверка", "operations": list(operations)})


def image_bytes(format="PNG", color="red"):
    output = io.BytesIO()
    Image.new("RGB", (80, 40), color).save(output, format=format)
    return output.getvalue()


def import_fixture(store):
    document = Document()
    p = document.add_paragraph()
    p.add_run("Привет, ").bold = True
    p.add_run("мир!").italic = True
    p.add_run().add_picture(io.BytesIO(image_bytes()), width=Cm(2))
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Товар 10"
    document.sections[0].header.paragraphs[0].text = "Колонтитул"
    output = io.BytesIO()
    document.save(output)
    store.import_document("пример.docx", output.getvalue())
    return output.getvalue()


def test_replacement_preserves_styles_picture_table_header_and_backup(store):
    original = import_fixture(store)
    snapshot = store.snapshot("пример.docx")
    image_id = snapshot["images"][0]["id"]
    image = store.image("пример.docx", image_id)
    result = store.edit("пример.docx", snapshot["revision"], plan(
        {"op": "replace_text", "target": "p0", "old": "Привет, мир", "new": "Добрый день"}))
    document = Document(store.path("пример.docx"))
    assert document.paragraphs[0].text == "Добрый день!"
    assert document.paragraphs[0].runs[0].bold
    assert document.paragraphs[0].runs[1].italic
    assert len(document.inline_shapes) == 1
    assert store.image("пример.docx", image_id) == image
    assert document.tables[0].cell(0, 0).text == "Товар 10"
    assert document.sections[0].header.paragraphs[0].text == "Колонтитул"
    restored = store.restore("пример.docx", result["backup_id"], result["revision"])
    assert store.path("пример.docx").read_bytes() == original
    assert restored["revision"] == sha256(original)
    assert len(store.backups("пример.docx")) == 2


def test_preview_never_writes(store):
    original = import_fixture(store)
    store.preview("пример.docx", sha256(original), plan(
        {"op": "replace_text", "target": "p0", "old": "мир", "new": "Word"}))
    assert store.path("пример.docx").read_bytes() == original
    assert store.backups("пример.docx") == []


def test_failed_transaction_never_commits_partial_changes(store):
    original = import_fixture(store)
    with pytest.raises(ValueError, match="was not found"):
        store.edit("пример.docx", sha256(original), plan(
            {"op": "replace_text", "target": "p0", "old": "мир", "new": "Word"},
            {"op": "replace_text", "target": "p0", "old": "НЕСУЩЕСТВУЮЩИЙ", "new": "error"}))
    assert store.path("пример.docx").read_bytes() == original


def test_stale_revision_rejected(store):
    original = import_fixture(store)
    store.edit("пример.docx", sha256(original), plan({"op": "add_paragraph", "text": "Новый текст"}))
    with pytest.raises(ValueError, match="changed"):
        store.edit("пример.docx", sha256(original), plan({"op": "add_paragraph", "text": "Устаревшая правка"}))


@pytest.mark.parametrize("name", ["../outside.docx", "C:/outside.docx", ".word_mcp/x.docx", "folder/../../x.docx", "x.docm", "\\\\server\\x.docx"])
def test_path_boundaries(store, name):
    with pytest.raises(ValueError):
        store.create(name)


def test_insertions_keep_original_ids_and_requested_order(store):
    revision = store.create("order.docx", "Заголовок")["revision"]
    store.edit("order.docx", revision, plan(
        {"op": "add_paragraph", "after": "p0", "text": "Первый"},
        {"op": "add_paragraph", "after": "p0", "text": "Второй"},
        {"op": "replace_text", "target": "p0", "old": "Заголовок", "new": "Новый заголовок"}))
    assert [p.text for p in Document(store.path("order.docx")).paragraphs] == ["Новый заголовок", "Первый", "Второй"]


def test_edit_table_cell_and_header(store):
    original = import_fixture(store)
    snapshot = store.snapshot("пример.docx")
    table_id = next(p["id"] for p in snapshot["paragraphs"] if p["text"] == "Товар 10")
    header_id = next(p["id"] for p in snapshot["paragraphs"] if p["text"] == "Колонтитул")
    store.edit("пример.docx", sha256(original), plan(
        {"op": "replace_text", "target": table_id, "old": "10", "new": "20"},
        {"op": "replace_text", "target": header_id, "old": "Колонтитул", "new": "Шапка"}))
    doc = Document(store.path("пример.docx"))
    assert doc.tables[0].cell(0, 0).text == "Товар 20"
    assert doc.sections[0].header.paragraphs[0].text == "Шапка"


def test_occurrence_selects_single_match(store):
    revision = store.create("repeat.docx", "слово слово слово")["revision"]
    store.edit("repeat.docx", revision, plan({"op": "replace_text", "target": "p0", "old": "слово", "new": "текст", "occurrence": 1}))
    assert Document(store.path("repeat.docx")).paragraphs[0].text == "слово текст слово"


def test_complex_fields_are_rejected(store):
    doc = Document()
    p = doc.add_paragraph("Текст")
    p._p.append(OxmlElement("w:fldSimple"))
    output = io.BytesIO()
    doc.save(output)
    result = store.import_document("field.docx", output.getvalue())
    with pytest.raises(ValueError, match="fields"):
        store.edit("field.docx", result["revision"], plan({"op": "replace_text", "target": "p0", "old": "Текст", "new": "Другой"}))


def test_add_and_replace_image_and_add_table(store):
    revision = store.create("images.docx")["revision"]
    asset = store.save_asset("red.png", image_bytes())
    result = store.edit("images.docx", revision, plan(
        {"op": "add_picture", "image_path": asset["path"], "alt_text": "Красный прямоугольник"},
        {"op": "add_table", "rows": [["Название", "Количество"], ["Товар", "10"]]}))
    snapshot = store.snapshot("images.docx")
    assert len(snapshot["images"]) == 1
    blue = store.save_asset("blue.png", image_bytes(color="blue"))
    store.edit("images.docx", result["revision"], plan({"op": "replace_picture", "image_id": snapshot["images"][0]["id"], "image_path": blue["path"]}))
    image = Image.open(io.BytesIO(store.image("images.docx", snapshot["images"][0]["id"])))
    assert image.getpixel((0, 0)) == (0, 0, 255)
    assert len(Document(store.path("images.docx")).tables) == 1


def test_shared_image_jpeg_keeps_correct_content_type(store):
    doc = Document()
    doc.add_picture(io.BytesIO(image_bytes("JPEG")), width=Cm(2))
    output = io.BytesIO()
    doc.save(output)
    revision = store.import_document("jpeg.docx", output.getvalue())["revision"]
    snapshot = store.snapshot("jpeg.docx")
    asset = store.save_asset("blue.png", image_bytes(color="blue"))
    store.edit("jpeg.docx", revision, plan({"op": "replace_picture", "image_id": snapshot["images"][0]["id"], "image_path": asset["path"]}))
    part = next(iter(store.image_parts(Document(store.path("jpeg.docx"))).values()))
    assert part.content_type == "image/jpeg"
    assert Image.open(io.BytesIO(part.blob)).format == "JPEG"


def test_bad_upload_and_zip_expansion_limits(store):
    with pytest.raises(ValueError):
        store.import_document("fake.docx", b"not a document")
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "x" * 1000)
    limited = DocumentStore(Settings(root=store.root, max_unpacked_bytes=100))
    with pytest.raises(ValueError, match="unpacked"):
        limited.import_document("bomb.docx", data.getvalue())


def test_multiple_process_lock_guard(store):
    path = store.path("lock.docx")
    another_store = DocumentStore(store.settings)
    with store.lock(path):
        with pytest.raises(ValueError, match="Another operation"):
            another_store.create("lock.docx")
    assert not list((store.internal / "locks").glob("*"))


def test_proposal_cannot_be_applied_twice(store):
    revision = store.create("proposal.docx", "Привет")["revision"]
    proposal = store.save_proposal("proposal.docx", revision, plan({"op": "replace_text", "target": "p0", "old": "Привет", "new": "Здравствуйте"}), "test")
    store.apply_proposal(proposal["proposal_id"])
    with pytest.raises(FileNotFoundError):
        store.apply_proposal(proposal["proposal_id"])
